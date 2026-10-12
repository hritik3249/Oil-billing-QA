#!/usr/bin/env python3
"""
Convert test cases CSV to Excel with formatting
Usage: python qa/create-excel.py
"""
import csv
import sys

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
except ImportError:
    print("openpyxl not installed. Installing now...")
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl"])
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# Create workbook
wb = Workbook()
ws = wb.active
ws.title = "Test Cases"

# Read CSV
with open('qa/test-cases.csv', 'r', encoding='utf-8') as f:
    reader = csv.reader(f)
    rows = list(reader)

# Write data
for row_idx, row in enumerate(rows, 1):
    for col_idx, value in enumerate(row, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=value)

        # Header formatting
        if row_idx == 1:
            cell.font = Font(bold=True, color="FFFFFF", size=11)
            cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        else:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

        # Add borders
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        cell.border = thin_border

# Set column widths
column_widths = {
    'A': 8,   # ID
    'B': 15,  # Module
    'C': 30,  # Title
    'D': 10,  # Priority
    'E': 35,  # Preconditions
    'F': 50,  # Steps
    'G': 40,  # Expected Result
    'H': 40,  # Actual Result
    'I': 10,  # Status
    'J': 30,  # Notes
}

for col, width in column_widths.items():
    ws.column_dimensions[col].width = width

# Set header row height
ws.row_dimensions[1].height = 30

# Freeze header row
ws.freeze_panes = 'A2'

# Save
wb.save('qa/test-cases.xlsx')
print(f"✓ Created qa/test-cases.xlsx with {len(rows)-1} test cases")
